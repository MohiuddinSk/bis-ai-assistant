"""Acquire and import public BIS directory metadata, never standards text.

The published-standards Excel URL is obtained with the public Download button
on the BIS Standards List. LIMS pages are followed through their own public
pagination. A failed acquisition leaves the previous manifest and datasets.

Install: python -m pip install -r scripts/requirements-directory.txt
Acquire: python scripts/sync_bis_directories.py acquire --standards-file path/to/export.xlsx --standards-url URL
Import:  python scripts/sync_bis_directories.py import
Check:   python scripts/sync_bis_directories.py validate
Refresh: python scripts/sync_bis_directories.py acquire --standards-file path/to/new-export.xlsx --standards-url URL --refresh
         python scripts/sync_bis_directories.py import
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urljoin, urlparse
from urllib.request import Request, urlopen

from lxml import html
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "data" / "reference" / "bis_directories"
RAW = DIRECTORY / "raw"
MANIFEST = DIRECTORY / "manifest.json"
PARSER_VERSION = "bis-directories-v5"
STANDARDS_LIST = "https://standards.bis.gov.in/website/published-standards/published-standards-list?selectedType=7&totalRow=1"
LAB_LIST = "https://lims.bis.gov.in/home/labs/"
SCOPE_SEARCH = "https://lims.bis.gov.in/home/search_is_number/"
# Chosen from actual BIS published-standard and LIMS results, not enumerated IDs.
SCOPE_IS_NUMBERS = ("4151", "15644", "210", "1417", "2062", "694", "13428", "456", "319", "15844")
MAX_BYTES = 20_000_000


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def atomic_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_bytes(content)
    os.replace(temporary, path)


def atomic_json(path: Path, value: object, *, compress: bool = False) -> None:
    data = (json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")
    atomic_bytes(path, gzip.compress(data, compresslevel=9) if compress else data)


def snapshot(url: str, *, refresh: bool, cache: dict[str, dict], min_delay: float) -> dict:
    if not refresh and url in cache:
        entry = cache[url]
        if (DIRECTORY / entry["path"]).exists():
            return entry
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in {"lims.bis.gov.in"}:
        raise ValueError(f"Unexpected acquisition host: {parsed.hostname}")
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            request = Request(url, headers={"User-Agent": "BIS-Bandhu-directory-metadata/1.0 (public pages; bounded refresh)"})
            with urlopen(request, timeout=35) as response:
                data = response.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES:
                raise ValueError("Source page exceeds the bounded download limit")
            if b"<html" not in data[:2000].lower():
                raise ValueError("Expected an HTML source page")
            digest = hashlib.sha256(data).hexdigest()
            relative = f"raw/{digest}.html.gz"
            path = DIRECTORY / relative
            if not path.exists():
                atomic_bytes(path, gzip.compress(data, compresslevel=9))
            time.sleep(min_delay)
            return {"url": url, "path": relative, "sha256": digest, "retrieved_at": utc_now(), "bytes": len(data)}
        except Exception as error:  # network and malformed responses retain last good manifest
            last_error = error
            time.sleep(min_delay * (attempt + 1))
    raise RuntimeError(f"Could not acquire public BIS page {url}: {last_error}")


def read_snapshot(entry: dict) -> bytes:
    path = DIRECTORY / entry["path"]
    content = gzip.decompress(path.read_bytes()) if path.suffix == ".gz" else path.read_bytes()
    if hashlib.sha256(content).hexdigest() != entry["sha256"]:
        raise ValueError(f"Source hash mismatch: {path}")
    return content


def next_page(document: html.HtmlElement, current_url: str) -> str | None:
    for anchor in document.xpath("//a"):
        if " ".join(anchor.itertext()).strip().lower() != "next":
            continue
        target = urljoin(current_url, anchor.get("href") or "")
        parsed = urlparse(target)
        parameters = parse_qs(parsed.query)
        if parsed.hostname == "lims.bis.gov.in" and "page" in parameters:
            # LIMS sometimes appends the next page to an existing page parameter.
            # The final value is the intended page; canonicalize it before caching.
            query = urlencode({key: values[-1] for key, values in parameters.items()})
            return parsed._replace(query=query).geturl()
    return None


def acquire(args: argparse.Namespace) -> None:
    if not args.standards_file or not args.standards_url:
        raise ValueError("Supply both --standards-file and the URL obtained from the official BIS Download button")
    source_path = Path(args.standards_file).resolve()
    source = source_path.read_bytes()
    suffix = source_path.suffix.lower()
    if suffix not in {".xlsx", ".csv", ".json"} or (suffix == ".xlsx" and source[:2] != b"PK"):
        raise ValueError("Supply an official XLSX, CSV or JSON metadata export")
    digest = hashlib.sha256(source).hexdigest()
    archived_path = RAW / f"{digest}{suffix}"
    if not archived_path.exists():
        atomic_bytes(archived_path, source)
    existing = json.loads(MANIFEST.read_text(encoding="utf-8")) if MANIFEST.exists() else {}
    cache = {entry["url"]: entry for entry in existing.get("sources", [])}
    sources: list[dict] = []
    for kind, first_url, limit in (
        ("laboratories", LAB_LIST, args.max_lab_pages),
        *((f"scopes:{number}", SCOPE_SEARCH + "?" + urlencode({"is_number__doc_no": number}), args.max_scope_pages)
          for number in SCOPE_IS_NUMBERS),
    ):
        url: str | None = first_url
        seen: set[str] = set()
        for _ in range(limit):
            if not url or url in seen:
                break
            seen.add(url)
            entry = snapshot(url, refresh=args.refresh, cache=cache, min_delay=args.delay)
            sources.append({**entry, "kind": kind})
            url = next_page(html.fromstring(read_snapshot(entry)), url)
        else:
            if url and url not in seen:
                print(f"Bounded at {limit} pages for {kind}; additional official pages were not fetched", file=sys.stderr)
    manifest = {
        "parser_version": PARSER_VERSION,
        "acquired_at": utc_now(),
        "standards": {
            "url": args.standards_url,
            "listing_url": STANDARDS_LIST,
            "path": str(archived_path.relative_to(ROOT)),
            "sha256": digest,
            "retrieved_at": utc_now(),
            "bytes": len(source),
        },
        "sources": sources,
    }
    atomic_json(MANIFEST, manifest)
    print(f"Acquired {len(sources)} public LIMS pages; manifest: {MANIFEST}")


def clean_text(element: html.HtmlElement) -> str:
    return " ".join(" ".join(element.itertext()).split())


def optional(value: object) -> str | None:
    text = " ".join(str(value or "").split())
    return text if text and text not in {"-", "--", "None", "N/A"} else None


def parse_standards(manifest: dict) -> tuple[list[dict], dict]:
    info = manifest["standards"]
    path = ROOT / info["path"]
    if hashlib.sha256(path.read_bytes()).hexdigest() != info["sha256"]:
        raise ValueError("Official Excel hash changed; acquire it again")
    if path.suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(value or "").strip() for value in next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[:6]]
        expected = ["Sl#", "Standard Number", "Date of Publish", "Title", "Type of Standard", "Degree of Equivalence"]
        if headers != expected:
            workbook.close()
            raise ValueError(f"Unexpected published-standards Excel columns: {headers}")
        rows = ((number, {key: value for key, value in zip(headers, row)})
                for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3))
    elif path.suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            table = list(csv.DictReader(stream))
        rows = ((number, row) for number, row in enumerate(table, start=2))
        workbook = None
    elif path.suffix == ".json":
        table = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(table, list) or not all(isinstance(row, dict) for row in table):
            raise ValueError("Official JSON export must be an array of metadata objects")
        rows = ((number, row) for number, row in enumerate(table, start=1))
        workbook = None
    else:
        raise ValueError("Unsupported standards export format")
    accepted: list[dict] = []
    rejected: list[dict] = []
    seen: set[str] = set()
    source_rows = 0
    for row_number, row in rows:
        source_rows += 1
        identifier = optional(row.get("Standard Number") or row.get("identifier"))
        published = optional(row.get("Date of Publish") or row.get("source_publication_date"))
        title = optional(row.get("Title") or row.get("title"))
        category = optional(row.get("Type of Standard") or row.get("category"))
        if not identifier or not title or not re.search(r"\d", identifier):
            rejected.append({"row": row_number, "reason": "missing identifier/title or no standard number"})
            continue
        key = re.sub(r"[^A-Z0-9]", "", identifier.upper())
        if key in seen:
            rejected.append({"row": row_number, "reason": "duplicate exact standard identity"})
            continue
        edition = re.search(r"(?::|\()\s*(\d{4})\)?$", identifier)
        part = re.search(r"\bPart\s*(\d+)\b", identifier, re.I)
        section = re.search(r"\bSec(?:tion)?\s*(\d+)\b", identifier, re.I)
        detail_url = optional(row.get("official_detail_url"))
        host = urlparse(detail_url).hostname if detail_url else None
        if detail_url and (urlparse(detail_url).scheme != "https" or not host or
                           (host != "bis.gov.in" and not host.endswith(".bis.gov.in"))):
            rejected.append({"row": row_number, "reason": "detail URL is not an official BIS HTTPS host"})
            continue
        seen.add(key)
        accepted.append({
            "identifier": identifier, "title": title, "category": category,
            "edition_year": edition.group(1) if edition else None,
            "part": part.group(1) if part else None,
            "section": section.group(1) if section else None,
            "status": None, "source_publication_date": published,
            "official_url": STANDARDS_LIST, "official_detail_url": detail_url,
            "retrieved_at": info["retrieved_at"],
            "provenance": f"BIS Published Standards {'Excel' if path.suffix == '.xlsx' else path.suffix[1:].upper()} export, row {row_number}",
            "source_row": row_number,
            "evidence_filename": None, "evidence_page": None,
        })
    if workbook:
        workbook.close()
    return accepted, {"source_rows": source_rows, "accepted": len(accepted), "rejected": len(rejected), "reasons": rejected[:100]}


def parse_laboratories(entries: list[dict]) -> tuple[list[dict], dict[str, dict], dict]:
    labs: dict[str, dict] = {}
    rejected: list[dict] = []
    count = 0
    for entry in entries:
        document = html.fromstring(read_snapshot(entry))
        tables = document.xpath("//table[@id='dataTable']")
        if len(tables) != 1:
            raise ValueError(f"No lab-directory table at {entry['url']}")
        for row in tables[0].xpath("./tbody/tr"):
            cells = row.xpath("./td")
            count += 1
            if len(cells) < 9:
                rejected.append({"source": entry["url"], "row": count, "reason": "missing lab columns"})
                continue
            code, name = optional(clean_text(cells[1])), optional(clean_text(cells[2]))
            if not code or not name:
                rejected.append({"source": entry["url"], "row": count, "reason": "missing code or name"})
                continue
            lines = [line.strip(" ,") for line in "".join(cells[3].itertext()).splitlines() if line.strip(" ,")]
            city = lines[-4] if len(lines) >= 5 else None
            state = lines[-2] if len(lines) >= 5 else None
            links = cells[8].xpath(".//a[@href]")
            detail = urljoin(entry["url"], links[0].get("href")) if links else None
            labs[code] = {
                "id": code, "code": code, "name": name,
                "location": ", ".join(part for part in (city, state) if part) or None,
                "state": state, "official_url": detail or entry["url"],
                "directory_valid_until": optional(clean_text(cells[7])),
                "recognition_status": None, "lab_type": None,
                "retrieved_at": entry["retrieved_at"],
                "provenance": f"BIS LIMS laboratory directory, {entry['url']}",
            }
    return list(labs.values()), labs, {"source_rows": count, "accepted": len(labs), "rejected": len(rejected), "reasons": rejected[:100]}


def parse_scopes(entries: list[dict], labs: dict[str, dict]) -> tuple[list[dict], dict]:
    scopes: list[dict] = []
    rejected: list[dict] = []
    seen: set[str] = set()
    count = 0
    for entry in entries:
        document = html.fromstring(read_snapshot(entry))
        tables = document.xpath("//table[@id='dataTable']")
        if len(tables) != 1:
            raise ValueError(f"No LIMS IS-search table at {entry['url']}")
        for row in tables[0].xpath("./tbody/tr"):
            cells = row.xpath("./td")
            count += 1
            if len(cells) < 9:
                rejected.append({"source": entry["url"], "row": count, "reason": "missing scope columns"})
                continue
            name, code, identifier = (optional(clean_text(cells[index])) for index in (1, 2, 3))
            if not name or not identifier:
                rejected.append({"source": entry["url"], "row": count, "reason": "missing lab or IS identifier"})
                continue
            lab_id = code if code else "name-" + hashlib.sha256(name.casefold().encode()).hexdigest()[:12]
            if lab_id not in labs:
                labs[lab_id] = {
                    "id": lab_id, "code": code, "name": name, "location": None, "state": None,
                    "official_url": entry["url"], "directory_valid_until": None,
                    "recognition_status": None, "lab_type": None,
                    "retrieved_at": entry["retrieved_at"],
                    "provenance": f"BIS LIMS IS-number search, {entry['url']}",
                }
            charge_lines = [" ".join(part.split()) for part in cells[6].xpath("./text()") if part.strip()]
            raw_charge = optional(charge_lines[0]) if charge_lines else None
            charge_note = optional(" ".join(charge_lines[1:]))
            # The top-level amount is a source-listed scope-row amount, not a
            # fabricated full-service quotation or a sum of clause charges.
            amount_match = re.fullmatch(r"(?:Rs\.?\s*)?(\d+(?:,\d{3})*(?:\.\d+)?)", raw_charge or "", re.I)
            charge = amount_match.group(1) if amount_match else None
            scope = {
                "laboratory_id": lab_id, "standard_identifier": identifier,
                "product_title": optional(clean_text(cells[4])),
                "grade_type": optional(clean_text(cells[5])),
                "testing_facility": None, "exclusions": None,
                # LIMS calls this merely "Validity Date" and often repeats the
                # lab-directory date. It is not identified as per-IS scope validity.
                "scope_valid_until": None,
                "search_validity_date": optional(clean_text(cells[7])),
                "charge_amount": charge, "charge_currency": "INR" if charge else None,
                "charge_basis": "BIS LIMS listed scope-row amount" if charge else None,
                "tax_treatment": "Excluding taxes" if charge else None,
                "charge_note": charge_note, "effective_date": None,
                "remarks": optional(" ".join(cells[8].xpath("./text()"))),
                "source_url": entry["url"], "retrieved_at": entry["retrieved_at"],
                "provenance": f"BIS LIMS IS-number search result, {entry['url']}",
            }
            key = json.dumps([scope[k] for k in (
                "laboratory_id", "standard_identifier", "product_title", "grade_type",
                "charge_amount", "charge_note", "search_validity_date", "remarks", "source_url",
            )])
            if key in seen:
                continue
            seen.add(key)
            scope["id"] = hashlib.sha256(key.encode()).hexdigest()[:20]
            scopes.append(scope)
    return scopes, {"source_rows": count, "accepted": len(scopes), "rejected": len(rejected), "reasons": rejected[:100]}


def import_data(_: argparse.Namespace) -> None:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    if manifest["parser_version"] != PARSER_VERSION:
        raise ValueError("Parser version differs from the source manifest")
    standards, standard_counts = parse_standards(manifest)
    directory_entries = [item for item in manifest["sources"] if item["kind"] == "laboratories"]
    scope_entries = [item for item in manifest["sources"] if item["kind"].startswith("scopes:")]
    laboratories, indexed, lab_counts = parse_laboratories(directory_entries)
    scopes, scope_counts = parse_scopes(scope_entries, indexed)
    laboratories = list(indexed.values())
    if len(standards) < 1000 or len(laboratories) < 100 or len(scopes) < 50:
        raise ValueError("Acquired data is unexpectedly small; prior datasets were not replaced")
    output = {
        "parser_version": PARSER_VERSION, "generated_at": utc_now(),
        "source_manifest": "manifest.json", "standards": standard_counts,
        "laboratories": lab_counts, "capabilities": scope_counts,
        "distinct_laboratories": len(laboratories),
        "distinct_capabilities": len(scopes),
        "capabilities_with_charge": sum(scope["charge_amount"] is not None for scope in scopes),
    }
    # Validate before any publication; atomic replacements preserve each last
    # good file if parsing/acquisition fails.
    atomic_json(DIRECTORY / "standards.json.gz", standards, compress=True)
    atomic_json(DIRECTORY / "laboratories.json.gz", {"laboratories": laboratories, "capabilities": scopes}, compress=True)
    atomic_json(DIRECTORY / "import_report.json", output)
    print(json.dumps(output, indent=2))


def validate(_: argparse.Namespace) -> None:
    report = json.loads((DIRECTORY / "import_report.json").read_text(encoding="utf-8"))
    standards = json.loads(gzip.decompress((DIRECTORY / "standards.json.gz").read_bytes()))
    labs = json.loads(gzip.decompress((DIRECTORY / "laboratories.json.gz").read_bytes()))
    assert len(standards) == report["standards"]["accepted"]
    assert len(labs["laboratories"]) == report["distinct_laboratories"]
    assert len(labs["capabilities"]) == report["distinct_capabilities"]
    assert len({row["id"] for row in labs["laboratories"]}) == len(labs["laboratories"])
    assert all(scope["laboratory_id"] in {lab["id"] for lab in labs["laboratories"]} for scope in labs["capabilities"])
    print(f"Valid: {len(standards)} standards, {len(labs['laboratories'])} labs, {len(labs['capabilities'])} capabilities")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    acquire_parser = commands.add_parser("acquire")
    acquire_parser.add_argument("--standards-file", type=Path, required=True)
    acquire_parser.add_argument("--standards-url", required=True)
    acquire_parser.add_argument("--refresh", action="store_true")
    acquire_parser.add_argument("--max-lab-pages", type=int, default=30)
    acquire_parser.add_argument("--max-scope-pages", type=int, default=3)
    acquire_parser.add_argument("--delay", type=float, default=1.0)
    commands.add_parser("import")
    commands.add_parser("validate")
    args = parser.parse_args()
    if args.command == "acquire":
        acquire(args)
    elif args.command == "import":
        import_data(args)
    else:
        validate(args)


if __name__ == "__main__":
    main()
